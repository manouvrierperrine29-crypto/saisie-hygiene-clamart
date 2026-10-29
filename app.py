import streamlit as st
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import io
import json
import base64
from datetime import date
from PIL import Image

# -----------------------------------------------------------------------------
# Configuration de la page Streamlit pour tablette / mobile
# -----------------------------------------------------------------------------
st.set_page_config(
    page_title="Saisie Terrain - Hygiène & Salubrité (Ville de Clamart)",
    page_icon="🏛️",
    layout="wide",
    initial_sidebar_state="expanded"
)

# -----------------------------------------------------------------------------
# Style CSS Institutionnel & Épuré (Bleu Marine #003366 & Gris Clair #F8F9FA)
# -----------------------------------------------------------------------------
st.markdown("""
    <style>
    /* Police globale et fond */
    html, body, [class*="css"] {
        font-family: 'Segoe UI', Roboto, Helvetica, Arial, sans-serif;
    }
    
    /* En-tête Institutionnel Ville de Clamart */
    .clamart-header {
        background: linear-gradient(135deg, #003366 0%, #002244 100%);
        color: white;
        padding: 20px 25px;
        border-radius: 12px;
        box-shadow: 0 4px 12px rgba(0,0,0,0.08);
        margin-bottom: 25px;
        display: flex;
        align-items: center;
        justify-content: space-between;
    }
    .clamart-title {
        font-size: 1.6rem;
        font-weight: 700;
        margin: 0;
        color: #FFFFFF;
        letter-spacing: 0.5px;
    }
    .clamart-subtitle {
        font-size: 0.95rem;
        color: #D1E3F8;
        margin-top: 4px;
        font-weight: 400;
    }
    .clamart-badge {
        background-color: rgba(255, 255, 255, 0.15);
        border: 1px solid rgba(255, 255, 255, 0.3);
        padding: 6px 14px;
        border-radius: 20px;
        font-size: 0.85rem;
        font-weight: 600;
        color: #F0F6FF;
    }

    /* Cartes de sections */
    .card-box {
        background-color: #FFFFFF;
        border-radius: 10px;
        padding: 20px;
        border: 1px solid #E2E8F0;
        border-left: 5px solid #003366;
        box-shadow: 0 2px 6px rgba(0,0,0,0.03);
        margin-bottom: 20px;
    }
    .card-title {
        color: #003366;
        font-size: 1.15rem;
        font-weight: 700;
        margin-bottom: 12px;
        display: flex;
        align-items: center;
        gap: 8px;
    }

    /* Boutons tactiles renforcés */
    .stButton>button {
        width: 100%;
        height: 3.2rem;
        font-weight: 600;
        font-size: 1.05rem;
        border-radius: 8px;
        background-color: #003366;
        color: white;
        border: none;
        transition: all 0.2s ease;
        box-shadow: 0 2px 4px rgba(0,51,102,0.2);
    }
    .stButton>button:hover {
        background-color: #002244;
        color: white;
        transform: translateY(-1px);
        box-shadow: 0 4px 8px rgba(0,51,102,0.3);
    }

    /* Onglets stylisés */
    .stTabs [data-baseweb="tab-list"] {
        gap: 8px;
        background-color: #F1F5F9;
        padding: 8px;
        border-radius: 10px;
    }
    .stTabs [data-baseweb="tab"] {
        height: 45px;
        border-radius: 6px;
        font-weight: 600;
        font-size: 0.95rem;
        color: #475569;
        background-color: transparent;
        border: none;
        padding: 0 16px;
    }
    .stTabs [aria-selected="true"] {
        background-color: #003366 !important;
        color: #FFFFFF !important;
        box-shadow: 0 2px 6px rgba(0,51,102,0.25);
    }

    /* Champs de saisie */
    .stTextInput input, .stSelectbox select, .stTextArea textarea, .stNumberInput input {
        border-radius: 6px !important;
        border: 1px solid #CBD5E1 !important;
        font-size: 0.98rem !important;
    }
    .stTextInput input:focus, .stTextArea textarea:focus {
        border-color: #003366 !important;
        box-shadow: 0 0 0 2px rgba(0,51,102,0.15) !important;
    }
    </style>
""", unsafe_allow_html=True)

# -----------------------------------------------------------------------------
# Bandeau En-Tête Officiel
# -----------------------------------------------------------------------------
st.markdown("""
    <div class="clamart-header">
        <div>
            <div class="clamart-title">🏛️ Ville de Clamart — Hygiène & Salubrité</div>
            <div class="clamart-subtitle">Direction Hygiène, Sécurité et Santé au Travail — Application Saisie Terrain (v4)</div>
        </div>
        <div class="clamart-badge">
            Inspection Habitat
        </div>
    </div>
""", unsafe_allow_html=True)

# -----------------------------------------------------------------------------
# Modèles types prédéfinis
# -----------------------------------------------------------------------------
PRESETS = {
    "Vierge": {},
    "Humidité & Ventilation défectueuse": {
        "type_logement": "Collectif",
        "typologie": "T3",
        "surface_mesuree": 52.0,
        "surface_declaree": 55.0,
        "hauteur_plafond": 2.50,
        "eclairage_naturel": "Moyen",
        "chauffage": "Gaz individuel",
        "loyer": 820,
        "mode_paiement": "Virement",
        "nb_occupants": 3,
        "nb_enfants": 1,
        "obs_cuisine": "Tags: Humidité / Infiltrations, Traces de moisissures. VMC obstruée par de la graisse et traces d'humidité sous évier.",
        "obs_sdb": "Tags: Traces de moisissures, Ventilation défectueuse/obstruée. Moisissures importantes sur les joints du bac à douche et au plafond.",
        "obs_ch1": "Tags: Humidité / Infiltrations. Traces de condensation sur le mur donnant sur l'extérieur.",
        "prescriptions_proprio": "• Remise en état du système de ventilation (VMC/Grilles)\n• Traitement des infiltrations d'eau en façade/toiture",
        "prescriptions_locataire": "• Nettoyage et désinfection des moisissures\n• Aération régulière du logement (2x10min/jour)",
        "conclusion": "Logement nécessitant des travaux d'aération et de traitement de l'humidité par le propriétaire, combiné à une meilleure aération quotidienne par l'occupant."
    },
    "Électricité vétuste & Risque de sécurité": {
        "type_logement": "Collectif",
        "typologie": "T2",
        "surface_mesuree": 38.0,
        "surface_declaree": 40.0,
        "hauteur_plafond": 2.40,
        "eclairage_naturel": "Satisfaisant",
        "chauffage": "Électricité",
        "loyer": 680,
        "mode_paiement": "Chèque",
        "nb_occupants": 2,
        "nb_enfants": 0,
        "obs_sejour": "Tags: Électricité vétuste/dangereuse. Prises de courant descellées du mur avec fils apparents.",
        "obs_cuisine": "Tags: Électricité vétuste/dangereuse. Absence de prise de terre à proximité du point d'eau.",
        "prescriptions_proprio": "• Remise aux normes de l'installation électrique\n• Remplacement des prises descellées et pose de disjoncteurs différentiels 30mA",
        "prescriptions_locataire": "• Dégagement des rallonges électriques inutilisées",
        "conclusion": "Installation électrique présentant un risque immédiat pour la sécurité des occupants. Mis en demeure du propriétaire pour mise aux normes."
    }
}

# -----------------------------------------------------------------------------
# Gestion des Modèles & Brouillons JSON (Barre Latérale)
# -----------------------------------------------------------------------------
with st.sidebar:
    st.markdown("### ⚙️ Modèles & Brouillons")
    preset_choice = st.selectbox("Charger un modèle prédéfini :", list(PRESETS.keys()))
    
    st.markdown("---")
    uploaded_json = st.file_uploader("📥 Importer un brouillon (.json)", type=["json"])
    if uploaded_json is not None:
        try:
            imported_data = json.load(uploaded_json)
            st.session_state['data_imported'] = imported_data
            st.success("✅ Brouillon restauré avec succès !")
        except Exception as e:
            st.error(f"Erreur de lecture du JSON : {e}")

# Fonction d'aide pour fond de cellule Word
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill_hex)
    shd.set(qn('w:val'), 'clear')
    tcPr.append(shd)

# -----------------------------------------------------------------------------
# Générateur du Rapport Word Officiel (.docx)
# -----------------------------------------------------------------------------
def create_docx_report(data, photos_list):
    doc = docx.Document()
    
    # Marges
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(10)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Entête
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_hdr = p_header.add_run(f"Clamart, le {data.get('date_visite_str', date.today().strftime('%d/%m/%Y'))}")
    r_hdr.bold = True
    r_hdr.font.size = Pt(10)

    p_service = doc.add_paragraph()
    r_srv = p_service.add_run("VILLE DE CLAMART\nDirection Hygiène Sécurité et Santé au Travail\n")
    r_srv.bold = True
    r_srv.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    r_srv.font.size = Pt(11)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("RAPPORT DE VISITE D'HYGIÈNE ET DE SALUBRITÉ DE L'HABITAT")
    r_title.bold = True
    r_title.font.size = Pt(13)
    r_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    p_title.paragraph_format.space_after = Pt(15)

    # Tableau Identification
    table_id = doc.add_table(rows=5, cols=2)
    table_id.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_id.autofit = False
    col_widths = [Inches(3.4), Inches(3.6)]
    
    hdr_cells = table_id.rows[0].cells
    hdr_cells[0].merge(hdr_cells[1])
    hdr_cells[0].text = "ÉLÉMENTS D'IDENTIFICATION"
    set_cell_background(hdr_cells[0], "003366")
    hdr_cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    hdr_cells[0].paragraphs[0].runs[0].font.bold = True
    
    row_data = [
        [f"Référence du rapport : {data.get('ref_rapport', '')}", f"Date de visite : {data.get('date_visite_str', '')}"],
        [f"Adresse du logement : {data.get('adresse', '')}", f"Réf. cadastrale / Année : {data.get('cadastre', '')}"],
        [f"Propriétaire(s) : {data.get('proprietaire', '')}\nCoord. : {data.get('coord_proprio', '')}", f"Mandataire : {data.get('mandataire', '')}\nCoord. : {data.get('coord_mandataire', '')}"],
        [f"Plaignant(s) : {data.get('plaignants', '')}", f"Date saisine Ville : {data.get('date_saisine', '')}"]
    ]
    
    for i, row in enumerate(row_data):
        cells = table_id.rows[i+1].cells
        for j in range(2):
            cells[j].text = row[j]
            cells[j].width = col_widths[j]
            set_cell_background(cells[j], "F2F4F7" if i % 2 == 0 else "FFFFFF")
            
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Descriptif logement
    p_desc = doc.add_paragraph()
    r_desc = p_desc.add_run("DESCRIPTIF DU LOGEMENT & CONSTAT")
    r_desc.bold = True
    r_desc.font.size = Pt(11)
    r_desc.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    table_desc = doc.add_table(rows=6, cols=2)
    table_desc.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_desc.autofit = False
    
    desc_rows = [
        [f"Type de logement : {data.get('type_logement', '')}", f"Typologie : {data.get('typologie', '')}"],
        [f"Surface mesurée : {data.get('surface_mesuree', '')} m²", f"Surface déclarée : {data.get('surface_declaree', '')} m²"],
        [f"Hauteur sous plafond : {data.get('hauteur_plafond', '')} m", f"Éclairage naturel : {data.get('eclairage_naturel', '')}"],
        [f"Type de chauffage : {data.get('chauffage', '')}", f"Date d'entrée : {data.get('date_entree', '')}"],
        [f"Loyer + charges : {data.get('loyer', '')} € ({data.get('mode_paiement', '')})", f"Occupants : {data.get('nb_occupants', '')} (dont {data.get('nb_enfants', '')} enfant(s))"],
        [f"Personnes présentes : {data.get('personnes_presentes', '')}", f"Conditions de visite : {data.get('conditions_visite', '')}"]
    ]
    
    for i, row in enumerate(desc_rows):
        cells = table_desc.rows[i].cells
        for j in range(2):
            cells[j].text = row[j]
            cells[j].width = col_widths[j]
            set_cell_background(cells[j], "FAFAFA")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Observations par pièce
    p_obs_head = doc.add_paragraph()
    r_obs = p_obs_head.add_run("OBSERVATIONS PAR PIÈCE")
    r_obs.bold = True
    r_obs.font.size = Pt(11)
    r_obs.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    pieces = [
        ("Cuisine", data.get("obs_cuisine", "")),
        ("Séjour", data.get("obs_sejour", "")),
        ("WC", data.get("obs_wc", "")),
        ("Salle de bain", data.get("obs_sdb", "")),
        ("Chambre 1", data.get("obs_ch1", "")),
        ("Chambre 2", data.get("obs_ch2", "")),
        ("Autres / Remarques complémentaires", data.get("obs_autres", ""))
    ]

    for nom_piece, obs_text in pieces:
        if obs_text and obs_text.strip():
            p_p = doc.add_paragraph()
            r_pname = p_p.add_run(f"• {nom_piece} : ")
            r_pname.bold = True
            r_pname.font.color.rgb = RGBColor(0x00, 0x56, 0xB3)
            p_p.add_run(obs_text)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Prescriptions
    p_avis = doc.add_paragraph()
    r_avis = p_avis.add_run("AVIS SANITAIRE & PRESCRIPTIONS")
    r_avis.bold = True
    r_avis.font.size = Pt(11)
    r_avis.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    table_presc = doc.add_table(rows=2, cols=1)
    table_presc.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    cell_prop = table_presc.rows[0].cells[0]
    cell_prop.width = Inches(7.0)
    p1 = cell_prop.paragraphs[0]
    r1 = p1.add_run("Prescriptions à charge du PROPRIÉTAIRE :\n")
    r1.bold = True
    p1.add_run(data.get("prescriptions_proprio", "Néant"))
    set_cell_background(cell_prop, "FDF2E9")

    cell_loc = table_presc.rows[1].cells[0]
    cell_loc.width = Inches(7.0)
    p2 = cell_loc.paragraphs[0]
    r2 = p2.add_run("Prescriptions à charge du LOCATAIRE :\n")
    r2.bold = True
    p2.add_run(data.get("prescriptions_locataire", "Néant"))
    set_cell_background(cell_loc, "EBF5FB")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Conclusion
    p_ccl = doc.add_paragraph()
    r_ccl_title = p_ccl.add_run("CONCLUSION GLOBALE :\n")
    r_ccl_title.bold = True
    r_ccl_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    p_ccl.add_run(data.get("conclusion", "Aucune observation particulière."))

    # Annexe Photos
    if photos_list:
        doc.add_page_break()
        p_photo_title = doc.add_paragraph()
        r_pt = p_photo_title.add_run("ANNEXE PHOTOGRAPHIQUE DE LA VISITE")
        r_pt.bold = True
        r_pt.font.size = Pt(13)
        r_pt.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        p_photo_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_photo_title.paragraph_format.space_after = Pt(15)

        for idx, (img_bytes, caption) in enumerate(photos_list, start=1):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                p_img.add_run().add_picture(io.BytesIO(img_bytes), width=Inches(4.5))
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_cap = p_cap.add_run(f"Photo {idx} : {caption if caption else 'Observation terrain'}")
                r_cap.font.italic = True
                r_cap.font.size = Pt(9)
                p_cap.paragraph_format.space_after = Pt(15)
            except Exception as img_err:
                st.warning(f"Note photo {idx} : {img_err}")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# -----------------------------------------------------------------------------
# Récupération des Données Initiales
# -----------------------------------------------------------------------------
init_data = PRESETS.get(preset_choice, {})
if 'data_imported' in st.session_state:
    init_data = st.session_state['data_imported']

# -----------------------------------------------------------------------------
# Navigation par Onglets Tactiles
# -----------------------------------------------------------------------------
tab1, tab2, tab3, tab4, tab5, tab6 = st.tabs([
    "📋 1. Identification", 
    "🏠 2. Logement", 
    "🔍 3. Constat / Pièces", 
    "📸 4. Photos Terrain",
    "⚖️ 5. Prescriptions",
    "📥 6. Exporter Word"
])

data = {}

with tab1:
    st.markdown('<div class="card-box"><div class="card-title">📋 Informations d\'Identification & Saisine</div>', unsafe_allow_html=True)
    col1, col2 = st.columns(2)
    with col1:
        data['ref_rapport'] = st.text_input("Référence du Rapport", value=init_data.get('ref_rapport', "REP-2026-001"))
        data['date_visite_str'] = st.date_input("Date de la visite", date.today()).strftime("%d/%m/%Y")
        data['adresse'] = st.text_input("Adresse du logement", value=init_data.get('adresse', ''), placeholder="Ex: 12 Rue de Paris, 92140 Clamart")
        data['cadastre'] = st.text_input("Réf. Cadastrale / Année construction", value=init_data.get('cadastre', ''), placeholder="Ex: AX 104 - 1972")
    with col2:
        data['proprietaire'] = st.text_input("Nom & Prénom Propriétaire(s)", value=init_data.get('proprietaire', ''))
        data['coord_proprio'] = st.text_input("Coordonnées Propriétaire", value=init_data.get('coord_proprio', ''))
        data['mandataire'] = st.text_input("Nom Mandataire (Gestionnaire/Syndic)", value=init_data.get('mandataire', ''))
        data['coord_mandataire'] = st.text_input("Coordonnées Mandataire", value=init_data.get('coord_mandataire', ''))
        data['plaignants'] = st.text_input("Nom & Prénom Plaignant(s)", value=init_data.get('plaignants', ''))
        data['date_saisine'] = st.text_input("Date saisine de la Ville", value=init_data.get('date_saisine', ''), placeholder="Ex: 10/10/2026")
    st.markdown('</div>', unsafe_allow_html=True)

with tab2:
    st.markdown('<div class="card-box"><div class="card-title">🏠 Caractéristiques du Logement</div>', unsafe_allow_html=True)
    col1, col2 = st.columns(2)
    with col1:
        data['type_logement'] = st.radio("Type de logement", ["Collectif", "Individuel"], index=0 if init_data.get('type_logement') == "Collectif" else 1, horizontal=True)
        data['typologie'] = st.selectbox("Typologie", ["T1", "T2", "T3", "T4", "T5", "T6", "Autre"], index=["T1", "T2", "T3", "T4", "T5", "T6", "Autre"].index(init_data.get('typologie', 'T3')))
        data['surface_mesuree'] = st.number_input("Surface mesurée (m²)", min_value=0.0, value=float(init_data.get('surface_mesuree', 45.0)), step=0.5)
        data['surface_declaree'] = st.number_input("Surface déclarée (m²)", min_value=0.0, value=float(init_data.get('surface_declaree', 48.0)), step=0.5)
        data['hauteur_plafond'] = st.number_input("Hauteur sous plafond (m)", min_value=0.0, value=float(init_data.get('hauteur_plafond', 2.50)), step=0.05)
        
        opts_eclairage = ["Insuffisant", "Moyen", "Satisfaisant", "Bon"]
        idx_ecl = opts_eclairage.index(init_data.get('eclairage_naturel', 'Satisfaisant')) if init_data.get('eclairage_naturel') in opts_eclairage else 2
        data['eclairage_naturel'] = st.select_slider("Éclairage naturel", options=opts_eclairage, value=opts_eclairage[idx_ecl])
    with col2:
        opts_chauff = ["Gaz individuel", "Gaz collectif", "Électricité", "Fioul / Autres"]
        idx_ch = opts_chauff.index(init_data.get('chauffage', 'Gaz individuel')) if init_data.get('chauffage') in opts_chauff else 0
        data['chauffage'] = st.selectbox("Type de chauffage", opts_chauff, index=idx_ch)
        
        data['date_entree'] = st.text_input("Date d'entrée dans le logement", value=init_data.get('date_entree', ''), placeholder="Ex: 01/05/2022")
        data['loyer'] = st.number_input("Loyer mensuel charges comprises (€)", min_value=0, value=int(init_data.get('loyer', 750)))
        data['mode_paiement'] = st.selectbox("Mode de paiement", ["Virement", "Chèque", "Prélèvement", "Espèces", "CAF direct"])
        data['nb_occupants'] = st.number_input("Nombre total d'occupants", min_value=1, value=int(init_data.get('nb_occupants', 3)))
        data['nb_enfants'] = st.number_input("Nombre d'enfants", min_value=0, value=int(init_data.get('nb_enfants', 1)))
        data['personnes_presentes'] = st.text_input("Personnes présentes lors de la visite", value=init_data.get('personnes_presentes', 'Occupant, Inspecteur Santé'))
        data['conditions_visite'] = st.selectbox("Conditions de visite", ["Normales", "Accès partiel", "Accès refusé", "Contestées"])
    st.markdown('</div>', unsafe_allow_html=True)

with tab3:
    st.markdown('<div class="card-box"><div class="card-title">🔍 Constats & Observations par pièce</div>', unsafe_allow_html=True)
    st.info("💡 **Astuce tablette :** Cochez les désordres fréquents ci-dessous pour pré-remplir les synthèses d'observation.")
    
    def render_piece_section(key_prefix, piece_name):
        st.markdown(f"##### 🚪 {piece_name}")
        quick_tags = st.multiselect(
            f"Désordres récurrents : {piece_name}",
            ["Humidité / Infiltrations", "Traces de moisissures", "Ventilation défectueuse/obstruée", "Problème d'étanchéité fenêtre", "Électricité vétuste/dangereuse", "Revêtement détérioré", "Manque de chauffage"],
            key=f"tags_{key_prefix}"
        )
        base_obs = init_data.get(f"obs_{key_prefix}", "")
        if quick_tags:
            tag_str = f"Tags: {', '.join(quick_tags)}. "
            if tag_str not in base_obs:
                base_obs = tag_str + base_obs
        txt_obs = st.text_area(f"Observations détaillées - {piece_name}", value=base_obs, key=f"txt_{key_prefix}", height=95)
        return txt_obs

    data['obs_cuisine'] = render_piece_section("cuisine", "Cuisine")
    data['obs_sejour'] = render_piece_section("sejour", "Séjour")
    data['obs_wc'] = render_piece_section("wc", "WC")
    data['obs_sdb'] = render_piece_section("sdb", "Salle de bain")
    data['obs_ch1'] = render_piece_section("ch1", "Chambre 1")
    data['obs_ch2'] = render_piece_section("ch2", "Chambre 2")
    data['obs_autres'] = st.text_area("Autres pièces / Remarques d'ensemble", value=init_data.get('obs_autres', ''), height=80)
    st.markdown('</div>', unsafe_allow_html=True)

with tab4:
    st.markdown('<div class="card-box"><div class="card-title">📸 Photos Terrain & Annexe Photographique</div>', unsafe_allow_html=True)
    st.caption("Prenez des clichés avec l'appareil photo de la tablette ou joignez des fichiers images.")
    
    col_cam, col_upload = st.columns(2)
    photos_captured = []
    
    with col_cam:
        st.markdown("##### 📷 Caméra en direct")
        camera_photo = st.camera_input("Prendre une photo")
        if camera_photo:
            cap_cam = st.text_input("Légende de la photo prise", value="Constat terrain")
            photos_captured.append((camera_photo.getvalue(), cap_cam))
            st.image(camera_photo, caption="Aperçu caméra", width=280)

    with col_upload:
        st.markdown("##### 🖼️ Fichiers images")
        uploaded_files = st.file_uploader("Sélectionner une ou plusieurs photos", type=["jpg", "jpeg", "png"], accept_multiple_files=True)
        if uploaded_files:
            for idx, file in enumerate(uploaded_files):
                cap_file = st.text_input(f"Légende pour {file.name}", value=f"Photo - {file.name}", key=f"cap_file_{idx}")
                photos_captured.append((file.getvalue(), cap_file))
                st.image(file, caption=f"Aperçu : {file.name}", width=230)
    st.markdown('</div>', unsafe_allow_html=True)

with tab5:
    st.markdown('<div class="card-box"><div class="card-title">⚖️ Prescriptions & Conclusion</div>', unsafe_allow_html=True)
    col1, col2 = st.columns(2)
    with col1:
        st.markdown("##### 🔨 À charge du PROPRIÉTAIRE")
        quick_prop = st.multiselect(
            "Prescriptions types (Propriétaire)",
            ["Remise en état du système de ventilation (VMC/Grilles)", "Traitement des infiltrations d'eau en façade/toiture", "Remise aux normes de l'installation électrique", "Changement du corps de chauffe / Radiateurs", "Réparation de la plomberie et fuites sous évier"]
        )
        base_p = init_data.get("prescriptions_proprio", "")
        if quick_prop:
            base_p += "\n• " + "\n• ".join(quick_prop)
        data['prescriptions_proprio'] = st.text_area("Prescriptions propriétaire détaillées", value=base_p, height=140)
        
    with col2:
        st.markdown("##### 🧹 À charge du LOCATAIRE")
        quick_loc = st.multiselect(
            "Prescriptions types (Locataire)",
            ["Nettoyage et désinfection des moisissures", "Aération régulière du logement (2x10min/jour)", "Entretien et désobstruction des bouches d'aération", "Dégagement des encombrants dans le logement"]
        )
        base_l = init_data.get("prescriptions_locataire", "")
        if quick_loc:
            base_l += "\n• " + "\n• ".join(quick_loc)
        data['prescriptions_locataire'] = st.text_area("Prescriptions locataire détaillées", value=base_l, height=140)

    st.markdown("##### 📝 Conclusion Globale du Rapport")
    data['conclusion'] = st.text_area("Synthèse générale d'inspection", value=init_data.get("conclusion", "Logement nécessitant des travaux d'aération et de remise en état. Notification transmise au propriétaire."), height=90)
    st.markdown('</div>', unsafe_allow_html=True)

with tab6:
    st.markdown('<div class="card-box"><div class="card-title">📥 Génération & Exports Officiels</div>', unsafe_allow_html=True)
    
    col_exp1, col_exp2 = st.columns(2)
    
    with col_exp1:
        st.markdown("### 📄 Rapport Word Officiel (.docx)")
        if photos_captured:
            st.success(f"📸 **{len(photos_captured)} photo(s)** prêtes à être intégrées dans l'annexe photographique.")
            
        docx_buffer = create_docx_report(data, photos_captured)
        file_name_docx = f"Rapport_Hygiene_{data.get('ref_rapport', 'Saisie')}.docx"
        
        st.download_button(
            label="📥 Télécharger le Rapport Officiel (.docx)",
            data=docx_buffer,
            file_name=file_name_docx,
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        
    with col_exp2:
        st.markdown("### 💾 Sauvegarde du Brouillon (.json)")
        st.caption("Générez le fichier JSON pour reprendre la saisie ultérieurement ou la sauvegarder.")
        
        # Préparation du dictionnaire exportable
        export_payload = data.copy()
        
        # Conversion du payload en octets UTF-8 pour un téléchargement binaire sans échec sur tablette
        json_str = json.dumps(export_payload, ensure_ascii=False, indent=2)
        json_bytes = json_str.encode('utf-8')
        file_name_json = f"Brouillon_{data.get('ref_rapport', 'Saisie')}.json"
        
        st.download_button(
            label="💾 Exporter le Brouillon (.json)",
            data=json_bytes,
            file_name=file_name_json,
            mime="application/json"
        )
    st.markdown('</div>', unsafe_allow_html=True)

st.markdown("---")
st.caption("Ville de Clamart — Direction Hygiène, Sécurité et Santé au Travail | Application officielle développée pour les inspecteurs de l'habitat.")

